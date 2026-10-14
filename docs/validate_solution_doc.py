#!/usr/bin/env python3
"""Validate the PoC guide and its generated Word export."""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
import zipfile
from pathlib import Path
from urllib.parse import urlparse

from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
EXPECTED_SECTIONS = [
    "1. What we built",
    "2. RAG in plain English",
    "3. What happens when we add documents",
    "4. The AWS architecture",
    "5. What happens when someone asks a question",
    "6. What the guardrails really mean",
    "7. How we would build the same demo in Azure",
    "8. AWS and Azure service mapping",
    "9. What changes when real Growthpoint documents arrive",
    "10. PoC versus production",
    "11. Sources",
]
REQUIRED_TERMS = (
    "25 fictional documents",
    "DEFAULT",
    "300 tokens",
    "Titan Text Embeddings V2",
    "1,024-dimension",
    "float32",
    "Amazon S3 Vectors",
    "cosine",
    "top five",
    "access_level",
    "ConverseStream",
    "Claude Haiku 4.5",
    "us.anthropic.claude-haiku-4-5-20251001-v1:0",
    "role dropdown is simulated",
    "text-embedding-3-small",
    "HNSW",
    "hybrid search",
    "Microsoft Entra ID",
    "Managed Identity",
    "Application Insights",
)
ALLOWED_LINK_HOSTS = {
    "docs.aws.amazon.com",
    "aws.amazon.com",
    "learn.microsoft.com",
    "teams.public.onecdn.static.microsoft",
}
FORBIDDEN = (
    "```mermaid",
    "Executive overview",
    "Solution architecture specification",
    "Operational considerations",
    "Claude 3 Haiku",
    "OpenSearch Serverless (this build)",
    "KNOWLEDGE_BASE_ID",
    "arn:aws:",
    "diagrams-mcp",
)


def words(text: str) -> int:
    return len(re.findall(r"\b[\w’'-]+\b", text))


def validate_markdown(path: Path) -> int:
    text = path.read_text(encoding="utf-8")
    count = words(text)
    if not 2500 <= count <= 3800:
        raise ValueError(f"Markdown word count {count} is outside 2,500–3,800")
    sections = re.findall(r"^## (.+)$", text, re.MULTILINE)
    if sections != EXPECTED_SECTIONS:
        raise ValueError(f"Unexpected section order: {sections}")
    image_refs = re.findall(r"!\[[^]]*]\(([^)]+)\)", text)
    expected_images = {
        "diagrams/aws-ingestion.png",
        "diagrams/aws-runtime.png",
        "diagrams/azure-ingestion.png",
        "diagrams/azure-runtime.png",
        "screenshots/01-landing.png",
        "screenshots/02-answer.png",
        "screenshots/03-sources.png",
    }
    if set(image_refs) != expected_images:
        raise ValueError(f"Expected exactly the four diagrams and three screenshots, got {image_refs}")
    for image in expected_images:
        if not (path.parent / image).is_file():
            raise FileNotFoundError(path.parent / image)
    lower = text.lower()
    for term in REQUIRED_TERMS:
        if term.lower() not in lower:
            raise ValueError(f"Required term missing: {term}")
    for term in FORBIDDEN:
        if term.lower() in lower:
            raise ValueError(f"Forbidden or stale content found: {term}")
    links = re.findall(r"https?://[^)\s]+", text)
    for link in links:
        host = urlparse(link).hostname
        if host not in ALLOWED_LINK_HOSTS:
            raise ValueError(f"Non-official link is not allowlisted: {link}")
    corpus = json.loads((ROOT / "src/mock/documents/corpus.json").read_text(encoding="utf-8"))
    if len(corpus.get("documents", [])) != 25:
        raise ValueError("Corpus cross-check failed: expected 25 documents")
    if re.search(r"<[a-zA-Z][^>]*>", text):
        raise ValueError("Raw HTML found in customer-facing Markdown")
    print(f"MARKDOWN_CHECK_OK: {count} words, 11 ordered sections, official links and 25-document corpus verified")
    return count


def document_text(document) -> str:
    chunks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def validate_docx(path: Path, markdown_words: int) -> None:
    from docx import Document

    if not zipfile.is_zipfile(path):
        raise ValueError(f"Not a valid OOXML ZIP package: {path}")
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        required_parts = {"word/document.xml", "word/styles.xml", "[Content_Types].xml"}
        if not required_parts <= names:
            raise ValueError(f"DOCX missing package parts: {required_parts - names}")
        document_xml = package.read("word/document.xml").decode("utf-8")
        styles_xml = package.read("word/styles.xml").decode("utf-8")
        footer_xml = "".join(
            package.read(name).decode("utf-8") for name in names if name.startswith("word/footer") and name.endswith(".xml")
        )
        media = [name for name in names if name.startswith("word/media/")]
        sizes = []
        for name in media:
            try:
                with Image.open(io.BytesIO(package.read(name))) as image:
                    sizes.append((name, image.size, image.width / image.height))
            except Exception:
                pass
    document = Document(path)
    text = document_text(document)
    lower = text.lower()
    headings = [p for p in document.paragraphs if p.style and p.style.name.startswith("Heading")]
    if len(headings) < 11:
        raise ValueError(f"DOCX has too few headings: {len(headings)}")
    if len(document.tables) < 3:
        raise ValueError(f"DOCX has too few tables: {len(document.tables)}")
    for section in EXPECTED_SECTIONS:
        if section.lower() not in lower:
            raise ValueError(f"DOCX missing section text: {section}")
    for forbidden in ("```mermaid", "<div", "<table", "knowledge_base_id", "arn:aws:"):
        if forbidden in lower:
            raise ValueError(f"DOCX contains forbidden literal content: {forbidden}")
    if "470635" not in styles_xml.upper():
        raise ValueError("Growthpoint aubergine heading colour is missing from Word styles")
    if "PAGE" not in footer_xml:
        raise ValueError("Page-number field is missing from footer")
    diagram_assets = [
        HERE / "diagrams/aws-ingestion.png",
        HERE / "diagrams/aws-runtime.png",
        HERE / "diagrams/azure-ingestion.png",
        HERE / "diagrams/azure-runtime.png",
    ]
    with zipfile.ZipFile(path) as package:
        media_hashes = {__import__("hashlib").sha256(package.read(name)).hexdigest() for name in media}
    for asset in diagram_assets:
        if __import__("hashlib").sha256(asset.read_bytes()).hexdigest() not in media_hashes:
            raise ValueError(f"Architecture image not embedded unchanged: {asset.name}")
    screenshot_assets = [
        HERE / "screenshots/01-landing.png",
        HERE / "screenshots/02-answer.png",
        HERE / "screenshots/03-sources.png",
    ]
    for asset in screenshot_assets:
        if __import__("hashlib").sha256(asset.read_bytes()).hexdigest() not in media_hashes:
            raise ValueError(f"Screenshot not embedded unchanged: {asset.name}")
    if len(sizes) < 8 or not any(ratio > 2.5 for _, _, ratio in sizes):
        raise ValueError(f"Expected logo + four diagrams + three screenshots; media={sizes}")
    estimated_pages = 1 + 1 + (markdown_words / 520) + 5
    if not 8 <= estimated_pages <= 15:
        print(f"WARNING: renderer-independent page estimate is {estimated_pages:.1f}, outside 8–15", file=sys.stderr)
    else:
        print(f"DOCX_CHECK_OK: valid OOXML, logo + 4 diagrams + 3 screenshots, {len(headings)} headings, {len(document.tables)} tables, page field; estimated {estimated_pages:.1f} pages")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--docx", type=Path)
    args = parser.parse_args()
    markdown_words = validate_markdown(args.markdown)
    if args.docx:
        validate_docx(args.docx, markdown_words)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
