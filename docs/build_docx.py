#!/usr/bin/env python3
"""Deterministically build the Growthpoint PoC guide as a branded DOCX."""

from __future__ import annotations

import base64
import copy
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
ROOT = DOCS.parent
SOURCE = DOCS / "SOLUTION-OVERVIEW.md"
BUILD = DOCS / "_build"
BUILD_MD = BUILD / "SOLUTION-OVERVIEW.word.md"
REFERENCE = BUILD / "growthpoint-reference.docx"
BODY_DOCX = BUILD / "solution-body.docx"
LOGO_SVG = ROOT / "assets" / "GrowthPoint-Logo (Global)-1.svg"
LOGO_PNG = BUILD / "growthpoint-logo.png"
OUTPUT = DOCS / "Growthpoint-Knowledge-Assistant-Solution-Overview.docx"
PANDOC = Path("/opt/homebrew/bin/pandoc")
AUBERGINE = "470635"
MAGENTA = "A4136C"
MUTED = "6F6872"


def set_fonts(style, name: str, size: float | None = None) -> None:
    style.font.name = name
    if size:
        style.font.size = Pt(size)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:cs"), "Calibri")


def style_reference() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    normal = doc.styles["Normal"]
    set_fonts(normal, "Aptos", 10.3)
    normal.font.color.rgb = RGBColor.from_string("2E2931")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for index, size in ((1, 18), (2, 14), (3, 11.5)):
        style = doc.styles[f"Heading {index}"]
        set_fonts(style, "Aptos Display", size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(AUBERGINE)
        style.paragraph_format.space_before = Pt(12 if index == 1 else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    set_fonts(doc.styles["Title"], "Aptos Display", 30)
    doc.styles["Title"].font.color.rgb = RGBColor.from_string(AUBERGINE)
    doc.styles["Title"].font.bold = True
    if "Caption" in doc.styles:
        caption = doc.styles["Caption"]
        set_fonts(caption, "Aptos", 8.5)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor.from_string(MUTED)
    doc.save(REFERENCE)


def render_logo() -> None:
    from playwright.sync_api import sync_playwright

    uri = "data:image/svg+xml;base64," + base64.b64encode(LOGO_SVG.read_bytes()).decode("ascii")
    markup = f'''<!doctype html><html><head><style>html,body{{margin:0;width:1200px;height:400px;overflow:hidden;background:white}}img{{width:1200px;height:400px;object-fit:contain}}</style></head><body><img id="logo" src="{uri}"></body></html>'''
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1200, "height": 400}, device_scale_factor=2)
        page.set_content(markup, wait_until="load")
        page.locator("#logo").screenshot(path=str(LOGO_PNG))
        browser.close()


def prepare_markdown() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    start = text.index("## 1. What we built")
    body = text[start:]
    body = re.sub(r"^## ", "# ", body, flags=re.MULTILINE)
    body = re.sub(r"^### ", "## ", body, flags=re.MULTILINE)
    BUILD_MD.write_text(body, encoding="utf-8")


def run(command: list[str], label: str) -> None:
    result = subprocess.run(command, cwd=ROOT, capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(f"{label} failed:\n{result.stdout}\n{result.stderr}")
    if result.stdout.strip():
        print(result.stdout.strip())


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def add_bottom_border(paragraph, color: str = MAGENTA) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_cover(doc: Document) -> None:
    first = doc._body._body[0]
    created = []

    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_p.paragraph_format.space_before = Pt(36)
    logo_p.add_run().add_picture(str(LOGO_PNG), width=Inches(5.4))
    created.append(logo_p._p)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(20)
    rule.paragraph_format.space_after = Pt(28)
    add_bottom_border(rule)
    created.append(rule._p)

    title = doc.add_paragraph(style="Title")
    title.add_run("Growthpoint Knowledge Assistant")
    title.paragraph_format.space_after = Pt(10)
    created.append(title._p)

    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("AWS RAG solution overview and Azure equivalent")
    subtitle_run.font.name = "Aptos Display"
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(24)
    created.append(subtitle._p)

    label = doc.add_paragraph()
    label_run = label.add_run("PROOF OF CONCEPT DEMO")
    label_run.bold = True
    label_run.font.name = "Aptos"
    label_run.font.size = Pt(10)
    label_run.font.color.rgb = RGBColor.from_string(MAGENTA)
    label.paragraph_format.space_after = Pt(8)
    created.append(label._p)

    status = doc.add_paragraph("Fictional demonstration data · Version 1.0 · October 2026")
    status.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    status.runs[0].font.size = Pt(9.5)
    created.append(status._p)

    break_p = doc.add_paragraph()
    break_p.add_run().add_break(WD_BREAK.PAGE)
    created.append(break_p._p)

    for element in created:
        first.addprevious(element)


def format_document() -> None:
    doc = Document(BODY_DOCX)
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.text = "Growthpoint Knowledge Assistant · Proof of concept"
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    add_bottom_border(header, AUBERGINE)
    add_page_number(section.footer.paragraphs[0])

    for table in doc.tables:
        table.style = "Light Shading Accent 1"
        table.autofit = True
        for index, cell in enumerate(table.rows[0].cells):
            set_cell_shading(cell, AUBERGINE)
            for paragraph in cell.paragraphs:
                for run_item in paragraph.runs:
                    run_item.font.color.rgb = RGBColor(255, 255, 255)
                    run_item.bold = True
                    run_item.font.name = "Aptos"
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = 1
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run_item in paragraph.runs:
                        run_item.font.name = "Aptos"
                        run_item.font.size = Pt(8.5)

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Figure "):
            paragraph.style = doc.styles["Caption"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True

    # At this point the only inline images are the four architecture diagrams
    # (the logo is added later on the cover page). They are wide, left-to-right
    # flows, so fit them to the text width and centre them.
    for shape in doc.inline_shapes:
        aspect = shape.width / shape.height
        if aspect > 1.6:
            shape.width = Inches(6.85)
            shape.height = int(shape.width / aspect)
            parent = shape._inline.getparent().getparent()
            # Center architecture figures without relying on their generated name.
            for paragraph in doc.paragraphs:
                if paragraph._p is parent:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = Pt(6)
                    paragraph.paragraph_format.space_after = Pt(4)
                    break

    add_cover(doc)
    doc.core_properties.title = "Growthpoint Knowledge Assistant — PoC Demo Guide"
    doc.core_properties.subject = "AWS RAG proof of concept and Azure equivalent architecture"
    doc.core_properties.author = "Growthpoint PoC team"
    doc.save(OUTPUT)


def main() -> int:
    BUILD.mkdir(parents=True, exist_ok=True)
    if not PANDOC.is_file():
        raise FileNotFoundError(PANDOC)
    run([sys.executable, "docs/validate_solution_doc.py", "--markdown", "docs/SOLUTION-OVERVIEW.md"], "Markdown validation")
    style_reference()
    render_logo()
    prepare_markdown()
    command = [
        str(PANDOC), str(BUILD_MD), "-o", str(BODY_DOCX),
        "--from=gfm", "--toc", "--toc-depth=2",
        f"--resource-path={DOCS}", f"--reference-doc={REFERENCE}",
        "--metadata=title:Growthpoint Knowledge Assistant",
    ]
    run(command, "Pandoc")
    format_document()
    run([sys.executable, "docs/validate_solution_doc.py", "--markdown", "docs/SOLUTION-OVERVIEW.md", "--docx", str(OUTPUT)], "DOCX validation")
    print(f"DOCX_BUILD_OK: {OUTPUT.relative_to(ROOT)} ({OUTPUT.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
